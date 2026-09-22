import importlib.util
import json
from pathlib import Path
import tempfile
import subprocess
import sys
import unittest
from unittest.mock import patch
from contextlib import contextmanager
from uuid import uuid4

ROOT = Path(__file__).resolve().parents[1]
TEST_ROOT = (ROOT / '.cache/rebuild/pipeline-python-tests').resolve()
assert TEST_ROOT.is_relative_to(ROOT)
TEST_ROOT.mkdir(parents=True, exist_ok=True)
tempfile.tempdir = str(TEST_ROOT)
@contextmanager
def fixture_directory():
    # Windows restricted execution cannot reopen mkdtemp's user-only ACL directories.
    # These tiny generated fixtures stay in the task's known writable test directory.
    path = (TEST_ROOT / uuid4().hex).resolve()
    assert path.is_relative_to(TEST_ROOT)
    path.mkdir()
    yield str(path)
def load(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module

reader_module = load('trusted_reader', ROOT / 'server' / 'sandbox-document-reader.py')

class ReaderTests(unittest.TestCase):
    def test_known_swig_import_deprecation_does_not_imply_document_loss(self):
        original = reader_module.Reader.read_text
        def harmless(reader, data):
            reader_module.warnings.warn('builtin type SwigPyObject has no __module__ attribute', DeprecationWarning)
            original(reader, data)
        def unsupported(reader, data):
            reader_module.warnings.warn('A source feature was altered', DeprecationWarning)
            original(reader, data)
        with fixture_directory() as directory:
            path = Path(directory) / 'source.txt'
            path.write_bytes(b'actual original')
            with patch.object(reader_module.Reader, 'read_text', harmless):
                result = reader_module.Reader(path, 'txt').read()
                self.assertTrue(result['coverage']['complete'])
                self.assertEqual(result['warnings'], [])
            with patch.object(reader_module.Reader, 'read_text', unsupported):
                result = reader_module.Reader(path, 'txt').read()
                self.assertFalse(result['coverage']['complete'])
                self.assertEqual(result['warnings'], ['Parser reported an unsupported or altered feature: DeprecationWarning'])

    def test_pdf_cold_parser_import_preserves_complete_physical_reading(self):
        import fitz
        with fixture_directory() as directory:
            path = Path(directory) / 'source.pdf'
            pdf = fitz.open()
            page = pdf.new_page()
            page.insert_text((40, 40), 'A complete physical page contains sufficient original digital text for inspection.')
            pdf.save(path)
            pdf.close()
            program = """import contextlib,io,json,runpy,sys
reader=runpy.run_path(sys.argv[1])['Reader']
with contextlib.redirect_stdout(io.StringIO()),contextlib.redirect_stderr(io.StringIO()):
    profile=reader(sys.argv[2],'pdf').read()
print(json.dumps({'complete':profile['coverage']['complete'],'pages':profile['coverage']['unitsRead'],'warnings':profile['warnings']}))
"""
            child = subprocess.run([sys.executable, '-c', program, str(ROOT / 'server/sandbox-document-reader.py'), str(path)], capture_output=True, text=True, check=True)
            self.assertEqual(json.loads(child.stdout), {'complete': True, 'pages': 1, 'warnings': []})

    def test_exact_text_boundaries(self):
        fixtures = json.loads((ROOT / 'architecture/contracts/source-serialization-cases.json').read_text(encoding='utf-8'))['readerTextBoundaryCases']
        with fixture_directory() as directory:
            path = Path(directory) / 'original.txt'
            for fixture in fixtures:
                count = fixture['inputRecipe']['bytes']['count']
                path.write_bytes(b'x' * count)
                result = reader_module.Reader(path, 'txt', Path(directory) / 'out.json').read()
                expected = fixture['expected']
                self.assertEqual(len(result['text']), expected['retainedChars'])
                self.assertEqual(result['coverage'], expected['coverage'])
                self.assertEqual(result['status'], expected['status'])
                self.assertEqual(result['warnings'], expected['warnings'])

    def test_sparse_workbook_formula_and_hidden_sheet(self):
        import openpyxl
        from openpyxl.styles import PatternFill
        with fixture_directory() as directory:
            path = Path(directory) / 'original.xlsx'
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet['B3'] = 'item'
            sheet['E9'] = '=1+1'
            sheet['XFD1048576'].fill = PatternFill('solid', fgColor='FFFFFF')
            sheet.row_dimensions[9].hidden = True
            hidden = workbook.create_sheet('hidden')
            hidden.sheet_state = 'hidden'
            hidden['A1'] = 'actual hidden source'
            workbook.save(path)
            result = reader_module.Reader(path, 'xlsx').read()
            self.assertEqual(result['coverage']['cellsRead'], 3)
            self.assertEqual(result['coverage']['missingFormulaCaches'], 1)
            formula = result['sheets'][0]['rows'][1]['cells'][0]
            self.assertTrue(formula['cacheMissing'])
            self.assertEqual(formula['formula'], '=1+1')
            self.assertEqual(result['sheets'][1]['state'], 'hidden')

    def test_inventory_sample_ties_and_sheet_101_empty(self):
        import openpyxl
        profile = load('criteria_profile', ROOT / 'server' / 'criteria-workbook-profile.py')
        workbook = openpyxl.Workbook()
        for row in range(1, 26):
            workbook.active.cell(row, 1, str(row))
        for index in range(100):
            workbook.create_sheet('S%d' % index)
        result = profile.inventory(workbook, workbook)
        samples = result['sheets'][0]['regions'][0]['sample']
        self.assertEqual([s['cell'] for s in samples], ['A%d' % i for i in list(range(1,13)) + list(range(14,26))])
        self.assertTrue(result['complete'])
        self.assertEqual(len(result['sheets']), 101)
        workbook.worksheets[100]['A1'] = 'not profiled'
        result = profile.inventory(workbook, workbook)
        self.assertFalse(result['complete'])
        self.assertEqual(result['sheets'][100]['profiledCells'], 0)

    def test_duplicate_detail_records_do_not_consume_unique_budget_twice(self):
        import openpyxl
        profile = load('criteria_details', ROOT / 'server' / 'criteria-workbook-profile.py')
        workbook = openpyxl.Workbook()
        for row in range(1, 51):
            workbook.active.cell(row, 1, 'x' * 4000)
        result = profile.details(workbook, workbook, {'ranges': [{'sheet': 'Sheet', 'range': 'A1:A50'}, {'sheet': 'Sheet', 'range': 'A1:B50'}]})
        self.assertEqual(result['cellsRead'], 50)
        self.assertTrue(result['complete'])
        self.assertEqual(sum(len(r['cells']) for r in result['ranges']), 100)
        self.assertGreater(len(json.dumps(result, ensure_ascii=False)), 220000)

    def test_requery_rejects_unknown_and_duplicate_keys(self):
        module = load('trusted_requery', ROOT / 'server' / 'document-requery.py')
        with fixture_directory() as directory:
            source, request = Path(directory) / 'source.txt', Path(directory) / 'requests.json'
            source.write_text('one\ntwo\nthree', encoding='utf-8')
            request.write_text('{"requests":[{"kind":"text","start":1,"end":2,"code":"ignored"}]}', encoding='utf-8')
            with self.assertRaises(ValueError):
                module.reread(source, 'txt', request)
            request.write_text('{"requests":[],"requests":[]}', encoding='utf-8')
            with self.assertRaises(ValueError):
                module.reread(source, 'txt', request)

    def test_docx_order_and_supplementary_headers(self):
        from docx import Document
        with fixture_directory() as directory:
            path = Path(directory) / 'original.docx'
            document = Document()
            document.add_paragraph('before table')
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = 'item'
            table.cell(0, 1).text = '30 이상'
            document.add_paragraph('after table')
            document.sections[0].header.paragraphs[0].text = 'header source'
            document.save(path)
            result = reader_module.Reader(path, 'docx').read()
            self.assertEqual([b['kind'] for b in result['blocks'][:3]], ['paragraph', 'table', 'paragraph'])
            self.assertTrue(any(b.get('text') == 'header source' for b in result['blocks']))
            self.assertTrue(all('page' not in b for b in result['blocks']))

    def test_image_metadata_never_invents_ocr(self):
        from PIL import Image
        with fixture_directory() as directory:
            path = Path(directory) / 'original.png'
            Image.new('RGB', (20, 20), 'white').save(path)
            result = reader_module.Reader(path, 'png').read()
            self.assertEqual(result['inventory']['pageCount'], 1)
            self.assertNotIn('text', result)
            self.assertEqual(result['images'], [])

    def test_pdf_inventory_uses_physical_pages_without_visual_30_page_limit(self):
        import fitz
        with fixture_directory() as directory:
            path = Path(directory) / 'original.pdf'
            pdf = fitz.open()
            for index in range(31):
                page = pdf.new_page()
                page.insert_text((40, 40), 'Actual physical source page %d has enough digital characters for text inspection.' % (index+1))
            pdf.save(path)
            pdf.close()
            result = reader_module.Reader(path, 'pdf').read()
            self.assertEqual(result['inventory']['pageCount'], 31)
            self.assertEqual(result['coverage']['unitsRead'], 31)
            self.assertEqual([page['page'] for page in result['pages']], list(range(1, 32)))

    def test_csv_logical_rows_and_text_separators_are_retained(self):
        with fixture_directory() as directory:
            path = Path(directory) / 'original.csv'
            path.write_bytes(b'ID,Value\r\n\r\nA,"two\nlines"\r\n')
            result = reader_module.Reader(path, 'csv').read()
            rows = result['sheets'][0]['rows']
            self.assertEqual([row['row'] for row in rows], [1, 2, 3])
            self.assertEqual(rows[1]['cells'], [])
            self.assertEqual(rows[2]['cells'][1]['value'], 'two\nlines')

if __name__ == '__main__':
    unittest.main()

"""Bounded trusted reader. Uploaded documents are data: no formula, macro or link execution."""
import csv
import datetime
import hashlib
import io
import json
import math
import os
import re
import sys
import time
import warnings
import zipfile

MAX_TEXT = 1500000
MAX_CELLS = 100000
MAX_IMAGES = 8
SENTINEL = '[계산 결과 없음: 수식 재계산 필요]'

def trace(stage, current, total=None):
    event = {'stage': stage, 'current': current}
    if total is not None:
        event['total'] = total
    print('TRACE_PROGRESS:' + json.dumps(event), flush=True)

def cell_value(value):
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)

def archive_check(data):
    archive = zipfile.ZipFile(io.BytesIO(data))
    entries = archive.infolist()
    if len(entries) > 20000 or sum(e.file_size for e in entries) > 100 * 1024 * 1024:
        raise ValueError('archive limit')
    names = set()
    for entry in entries:
        if entry.filename in names or entry.flag_bits & 1 or entry.compress_type not in (0, 8):
            raise ValueError('archive structure')
        names.add(entry.filename)
    return archive

class Reader:
    def __init__(self, path=None, kind=None, output=None):
        self.path, self.kind = path, kind
        self.started = time.monotonic()
        self.profile = {'kind': kind, 'sha256': '', 'status': 'ready', 'inventory': {}, 'coverage': {
            'complete': True, 'unitsTotal': 0, 'unitsRead': 0, 'cellsTotal': 0, 'cellsRead': 0,
            'textChars': 0, 'truncated': False, 'truncatedReasons': []}, 'warnings': [], 'images': []}
        self.result = self.profile
        self.coverage = self.profile['coverage']

    def warn(self, message):
        message = str(message)[:1000]
        if message not in self.profile['warnings'] and len(self.profile['warnings']) < 100:
            self.profile['warnings'].append(message)

    def partial(self, reason):
        self.coverage['complete'] = False
        self.coverage['truncated'] = True
        if reason not in self.coverage['truncatedReasons']:
            self.coverage['truncatedReasons'].append(reason)
        self.warn(reason)

    def text(self, value):
        value = '' if value is None else str(value)
        remaining = max(0, MAX_TEXT - self.coverage['textChars'])
        if len(value) > remaining:
            self.partial('Text extraction limit (1500000 characters) reached.')
            value = value[:remaining]
        self.coverage['textChars'] += len(value)
        return value

    def available(self):
        if time.monotonic() - self.started > 150:
            self.partial('Reader elapsed time limit reached')
            return False
        if self.coverage['cellsRead'] >= MAX_CELLS:
            self.partial('Cell limit reached')
            return False
        if self.coverage['textChars'] >= MAX_TEXT:
            self.partial('Text character limit reached')
            return False
        return True

    def save_image(self, data, **metadata):
        if len(self.profile['images']) >= MAX_IMAGES:
            self.partial('Image count limit reached')
            return None
        try:
            from PIL import Image
            image = Image.open(io.BytesIO(data))
            if image.width * image.height > 40000000:
                raise ValueError('pixel limit')
            image.seek(0)
            image.thumbnail((3000, 3000))
            if image.mode not in ('RGB', 'RGBA'):
                image = image.convert('RGBA')
            path = '/home/user/document-image-%d.png' % len(self.profile['images'])
            image.save(path, 'PNG')
            info = {'path': path, 'mime': 'image/png', 'width': image.width, 'height': image.height, **metadata}
            self.profile['images'].append(info)
            return info
        except Exception:
            self.coverage['complete'] = False
            self.warn('Embedded image could not be decoded within supported limits')
            return None

    def read(self, path=None, kind=None):
        self.path, self.kind = path or self.path, kind or self.kind
        self.profile['kind'] = self.kind
        size = os.path.getsize(self.path)
        if size < 1:
            raise ValueError('source size')
        with open(self.path, 'rb') as handle:
            data = handle.read()
        digest = hashlib.sha256(data).hexdigest()
        self.profile['sha256'] = digest
        self.profile['inventory'].update(sourceBytes=len(data), bytesRead=len(data), sha256=digest)
        trace('file', 1, 1)
        handlers = {'xlsx': self.read_xlsx, 'pdf': self.read_pdf, 'docx': self.read_docx,
                    'csv': self.read_csv, 'txt': self.read_text, 'md': self.read_text, 'json': self.read_text,
                    'png': self.read_image, 'jpg': self.read_image, 'jpeg': self.read_image, 'webp': self.read_image}
        if self.kind not in handlers:
            raise ValueError('unsupported kind')
        with warnings.catch_warnings(record=True) as observed:
            warnings.simplefilter('always')
            handlers[self.kind](data)
            for warning in observed:
                if issubclass(warning.category, DeprecationWarning) and re.fullmatch(
                        r'builtin type (?:SwigPyPacked|SwigPyObject|swigvarlink) has no __module__ attribute',
                        str(warning.message)):
                    continue
                self.partial('Parser reported an unsupported or altered feature: ' + warning.category.__name__)
        self.profile['status'] = 'ready' if self.coverage['complete'] else 'partial'
        return self.profile

    def read_text(self, data):
        decoded = None
        for encoding in ('utf-8-sig', 'utf-16', 'cp949'):
            try:
                decoded = data.decode(encoding)
                self.profile['inventory']['encoding'] = encoding
                break
            except UnicodeError:
                pass
        if decoded is None:
            raise ValueError('encoding')
        self.profile['text'] = self.text(decoded)
        self.coverage.update(unitsTotal=1, unitsRead=1)

    def read_csv(self, data):
        from openpyxl.utils import get_column_letter
        decoded = data.decode('utf-8-sig')
        try:
            dialect = csv.Sniffer().sniff(decoded[:65536], delimiters=',;\t|')
        except csv.Error:
            dialect = csv.excel
            self.warn('CSV delimiter detection failed; using comma')
        csv.field_size_limit(1500000)
        rows = []
        complete = True
        count = 0
        max_column = 0
        for number, values in enumerate(csv.reader(io.StringIO(decoded, newline=''), dialect), 1):
            count = number
            if number > 100000 or not self.available():
                self.partial('CSV row or content limit reached')
                complete = False
                break
            cells = []
            for column, value in enumerate(values, 1):
                if not self.available():
                    complete = False
                    break
                cells.append({'cell': get_column_letter(column) + str(number), 'column': column,
                              'columnName': get_column_letter(column), 'value': self.text(value)})
                self.coverage['cellsRead'] += 1
            rows.append({'row': number, 'cells': cells})
            max_column = max(max_column, len(values))
            self.coverage['unitsRead'] += 1
            self.coverage['cellsTotal'] += len(values)
            if not complete:
                break
        self.coverage['unitsTotal'] = count if complete else None
        self.coverage['rowsRead'] = len(rows)
        self.profile['inventory'].update(rowCount=count if complete else None, columnCount=max_column)
        if not complete:
            self.profile['inventory']['rowCountAtLeast'] = count
        self.profile['sheets'] = [{'name': 'CSV', 'state': 'visible', 'maxRow': len(rows), 'maxColumn': max_column,
                                  'mergedRanges': [], 'hiddenRows': [], 'hiddenColumns': [], 'rows': rows, 'complete': complete}]

    def read_xlsx(self, data):
        import openpyxl
        from openpyxl.cell.cell import MergedCell
        from openpyxl.utils import get_column_letter
        archive_check(data)
        workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=False, read_only=False, keep_links=False)
        cached = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=False, keep_links=False)
        self.profile['sheets'] = []
        inventory = []
        populated = {}
        for sheet in workbook.worksheets:
            cells = sorted((c for c in sheet._cells.values() if not isinstance(c, MergedCell) and (c.value is not None or c.comment)), key=lambda c: (c.row, c.column))
            populated[sheet.title] = cells
            inventory.append({'name': sheet.title, 'state': sheet.sheet_state, 'maxRow': sheet.max_row,
                              'maxColumn': sheet.max_column, 'storedCells': len(cells), 'mergedRanges': len(sheet.merged_cells.ranges),
                              'images': len(sheet._images), 'charts': len(sheet._charts), 'tables': list(sheet.tables), 'protected': bool(sheet.protection.sheet)})
        names = list(workbook.defined_names.items())
        self.profile['inventory'].update(sheetCount=len(inventory), sheets=inventory,
            definedNames=[{'name': name, 'value': str(value.attr_text)} for name, value in names[:500]])
        if len(names) > 500:
            self.partial('Defined name limit reached')
        self.coverage['unitsTotal'] = len(inventory)
        self.coverage['cellsTotal'] = sum(len(v) for v in populated.values())
        self.coverage['missingFormulaCaches'] = 0
        total_rows = 0
        for sheet_index, sheet in enumerate(workbook.worksheets):
            if sheet_index >= 100:
                self.partial('Worksheet limit reached')
                break
            cached_sheet = cached[sheet.title]
            output = {'name': sheet.title, 'state': sheet.sheet_state, 'maxRow': sheet.max_row, 'maxColumn': sheet.max_column,
                      'mergedRanges': [str(r) for r in sheet.merged_cells.ranges],
                      'hiddenRows': [i for i, dim in sheet.row_dimensions.items() if dim.hidden],
                      'hiddenColumns': [{'column': key, 'min': dim.min, 'max': dim.max} for key, dim in sheet.column_dimensions.items() if dim.hidden],
                      'autoFilter': sheet.auto_filter.ref,
                      'tables': [{'name': table.name, 'ref': table.ref} for table in sheet.tables.values()],
                      'headersFooters': {'header': str(sheet.oddHeader), 'footer': str(sheet.oddFooter)},
                      'imageInventory': [], 'chartInventory': [], 'rows': [], 'complete': True}
            current_row = None
            for cell in populated[sheet.title]:
                if not self.available() or total_rows >= 100000:
                    self.partial('Workbook cell, row, text or time limit reached')
                    output['complete'] = False
                    break
                if current_row != cell.row:
                    output['rows'].append({'row': cell.row, 'cells': []})
                    current_row = cell.row
                    total_rows += 1
                value = cell_value(cell.value)
                record = {'cell': cell.coordinate, 'column': cell.column, 'columnName': get_column_letter(cell.column), 'value': value}
                if isinstance(value, str):
                    record['value'] = self.text(value)
                if cell.data_type == 'f':
                    record['formula'] = record['value']
                    record['cachedValue'] = cell_value(cached_sheet[cell.coordinate].value)
                    if record['cachedValue'] is None or record['cachedValue'] == '':
                        record['cacheMissing'] = True
                        self.coverage['missingFormulaCaches'] += 1
                if cell.comment:
                    record['comment'] = self.text(cell.comment.text)
                if cell.hyperlink:
                    record['hyperlink'] = self.text(cell.hyperlink.target or cell.hyperlink.location or '')
                if cell.number_format:
                    record['numberFormat'] = cell.number_format
                output['rows'][-1]['cells'].append(record)
                self.coverage['cellsRead'] += 1
            for img in sheet._images:
                anchor = getattr(img.anchor, '_from', None)
                location = get_column_letter(anchor.col + 1) + str(anchor.row + 1) if anchor is not None else str(img.anchor)
                output['imageInventory'].append({'cell': location, 'width': img.width, 'height': img.height})
                self.save_image(img._data(), sheet=sheet.title, cell=location)
            for chart in sheet._charts:
                output['chartInventory'].append({'type': type(chart).__name__})
                self.partial('Workbook chart visual was not read')
            self.profile['sheets'].append(output)
            self.coverage['unitsRead'] += 1
            trace('sheet', self.coverage['unitsRead'], self.coverage['unitsTotal'])
        if self.coverage['missingFormulaCaches']:
            self.warn('Formula saved results are missing; formulas were not recalculated')
        self.coverage['rowsRead'] = total_rows
        workbook.close()
        cached.close()

    def read_pdf(self, data):
        import fitz
        pdf = fitz.open(stream=data, filetype='pdf')
        if pdf.needs_pass:
            raise ValueError('encrypted PDF')
        self.profile['inventory'].update(pageCount=len(pdf), metadata=pdf.metadata, embeddedFiles=pdf.embfile_count())
        self.coverage['unitsTotal'] = len(pdf)
        self.profile['pages'] = []
        for index, page in enumerate(pdf):
            if index >= 300 or not self.available():
                self.partial('PDF page or content limit reached')
                break
            text = self.text(page.get_text('text', sort=True))
            blocks = [{'bbox': list(block[:4]), 'text': self.text(block[4]), 'type': 'text'} for block in page.get_text('blocks', sort=True) if block[6] == 0]
            tables = []
            try:
                for table in page.find_tables().tables:
                    rows = [[self.text(value) for value in row] for row in table.extract()]
                    tables.append({'bbox': list(table.bbox), 'header': {'names': table.header.names, 'external': table.header.external}, 'rows': rows})
            except Exception:
                self.warn('PDF table detection could not inspect a page')
            fonts = page.get_fonts(full=True)
            image_inventory = [{'xref': item[0], 'width': item[2], 'height': item[3]} for item in page.get_images(full=True)]
            self.profile['pages'].append({'page': index + 1, 'width': page.rect.width, 'height': page.rect.height,
                'rotation': page.rotation, 'text': text, 'blocks': blocks, 'tables': tables, 'imageInventory': image_inventory, 'fonts': fonts, 'complete': True})
            self.coverage['unitsRead'] += 1
            trace('page', self.coverage['unitsRead'], self.coverage['unitsTotal'])
            if len(''.join(text.split())) < 40 or any(item[1] == 'n/a' and ('CID' in item[2] or 'CJK' in item[3]) for item in fonts):
                scale = min(1.7, 2800 / max(page.rect.width, page.rect.height))
                self.save_image(page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False).tobytes('png'), page=index + 1, sourceKind='pdf-page')
        pdf.close()

    def read_image(self, data):
        from PIL import Image
        image = Image.open(io.BytesIO(data))
        if image.width * image.height > 40000000:
            raise ValueError('pixel limit')
        frames = getattr(image, 'n_frames', 1)
        self.profile['inventory'].update(width=image.width, height=image.height, format=image.format, frames=frames, pageCount=1)
        self.coverage.update(unitsTotal=frames, unitsRead=1)
        if frames > 1:
            self.partial('Only the first image frame was inspected')

    def read_docx(self, data):
        from docx import Document
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        from lxml import etree
        archive = archive_check(data)
        document = Document(io.BytesIO(data))
        self.profile['blocks'] = []
        self.profile['sections'] = [{'section': i+1, 'pageWidth': section.page_width, 'pageHeight': section.page_height, 'orientation': str(section.orientation)} for i, section in enumerate(document.sections)]
        self.profile['inventory']['coreProperties'] = {key: cell_value(getattr(document.core_properties, key, None)) for key in ('title', 'subject', 'author', 'keywords', 'created', 'modified')}
        def table_data(table, depth=0):
            if depth > 8:
                self.partial('Nested Word table depth limit reached')
                return [], []
            rows, nested = [], []
            for r, row in enumerate(table.rows, 1):
                values = []
                for c, cell in enumerate(row.cells, 1):
                    if not self.available():
                        return rows, nested
                    values.append(self.text(cell.text))
                    self.coverage['cellsRead'] += 1
                    self.coverage['cellsTotal'] += 1
                    for n, child in enumerate(cell.tables, 1):
                        child_rows, grandchildren = table_data(child, depth+1)
                        nested.append({'row': r, 'column': c, 'index': n, 'rows': child_rows, 'nestedTables': grandchildren})
                rows.append(values)
            return rows, nested
        sources = [('body', document.element.body, document)]
        seen_parts = set()
        for i, section in enumerate(document.sections, 1):
            for kind in ('header', 'footer', 'first_page_header', 'first_page_footer', 'even_page_header', 'even_page_footer'):
                part = getattr(section, kind)
                if part.is_linked_to_previous or str(part.part.partname) in seen_parts:
                    continue
                seen_parts.add(str(part.part.partname))
                sources.append(('section%d/%s' % (i, kind), part._element, part))
        total = sum(sum(child.tag in (qn('w:p'), qn('w:tbl')) for child in element) for _, element, _ in sources)
        self.coverage['unitsTotal'] = total
        for source, element, owner in sources:
            for child in element:
                if child.tag not in (qn('w:p'), qn('w:tbl')):
                    continue
                if not self.available():
                    break
                block = {'index': len(self.profile['blocks'])+1, 'source': source}
                if child.tag == qn('w:p'):
                    paragraph = Paragraph(child, owner)
                    block.update(kind='paragraph', text=self.text(paragraph.text), style=paragraph.style.name if paragraph.style else '')
                else:
                    rows, nested = table_data(Table(child, owner))
                    block.update(kind='table', rows=rows, nestedTables=nested, complete=self.coverage['complete'])
                self.profile['blocks'].append(block)
                self.coverage['unitsRead'] += 1
                if self.coverage['unitsRead'] % 100 == 0 or self.coverage['unitsRead'] == total:
                    trace('blocks', self.coverage['unitsRead'], total)
        self.profile['supplementaryText'] = []
        parser = etree.XMLParser(resolve_entities=False, no_network=True)
        for name in archive.namelist():
            if name in ('word/footnotes.xml', 'word/endnotes.xml', 'word/comments.xml'):
                element = etree.fromstring(archive.read(name), parser)
                self.profile['supplementaryText'].append({'part': name, 'text': self.text('\n'.join(element.xpath('//*[local-name()="t"]/text()')))})
            if name.startswith('word/media/'):
                self.save_image(archive.read(name), description='Word embedded image')
        root = etree.fromstring(archive.read('word/document.xml'), parser)
        for index, textbox in enumerate(root.xpath('//*[local-name()="txbxContent"]'), 1):
            self.profile['supplementaryText'].append({'part': 'textbox%d' % index, 'text': self.text('\n'.join(textbox.xpath('.//*[local-name()="t"]/text()')))})

def main():
    reader = Reader(sys.argv[1], sys.argv[2])
    try:
        reader.read()
    except Exception as error:
        reader.profile['status'] = 'unsupported'
        reader.coverage['complete'] = False
        reader.warn('Unsupported document or parser failure: ' + type(error).__name__)
    with open(sys.argv[3], 'w', encoding='utf-8') as handle:
        json.dump(reader.profile, handle, ensure_ascii=False, allow_nan=False)
    print(json.dumps({'cellsRead': reader.coverage['cellsRead'], 'unitsRead': reader.coverage['unitsRead'], 'unitsTotal': reader.coverage['unitsTotal']}))

if __name__ == '__main__':
    main()

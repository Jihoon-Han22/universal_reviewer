"""Sparse workbook discovery inventory and bounded exact-cell reads."""
import datetime
import io
import json
import sys
import zipfile
import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter, range_boundaries

SENTINEL = '[계산 결과 없음: 수식 재계산 필요]'

def text(value):
    if value is None:
        return ''
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    return str(value)

def stored_cells(sheet):
    return sorted((cell for cell in sheet._cells.values() if not isinstance(cell, MergedCell) and (cell.value is not None or cell.comment)), key=lambda cell: (cell.row, cell.column))

def record(cell, cached_sheet, sheet):
    formula = cell.data_type == 'f'
    cached_value = cached_sheet[cell.coordinate].value if formula else cell.value
    uncached = formula and (cached_value is None or cached_value == '')
    value = SENTINEL if uncached else text(cached_value)
    comment = cell.comment.text if cell.comment else ''
    hidden_column = any(dimension.hidden and (dimension.min or 1) <= cell.column <= (dimension.max or dimension.min or 1) for dimension in sheet.column_dimensions.values())
    result = {'cell': cell.coordinate, 'text': value[:4000], 'styleId': cell.style_id,
              'hiddenRow': bool(sheet.row_dimensions.get(cell.row) and sheet.row_dimensions[cell.row].hidden), 'hiddenColumn': hidden_column}
    for merged in sheet.merged_cells.ranges:
        if cell.coordinate in merged:
            result['mergedRange'] = str(merged)
            break
    if formula:
        result['formula'] = text(cell.value)[:1000]
    if uncached:
        result['uncachedFormula'] = True
    if comment:
        result['comment'] = comment[:4000]
    if len(value) > 4000 or len(comment) > 4000:
        result['truncated'] = True
    return result

def bands(values):
    groups = []
    for value in sorted(set(values)):
        if not groups or value != groups[-1][-1] + 1:
            groups.append([])
        groups[-1].append(value)
    return groups

def inventory(workbook, cached):
    result = {'sheets': [], 'warnings': [], 'complete': True, 'limits': {'sheets': 100, 'profiledCells': 200000, 'regions': 300}}
    remaining, region_count = 200000, 0
    for index, sheet in enumerate(workbook.worksheets):
        population = stored_cells(sheet)
        selected = population[:remaining] if index < 100 else []
        remaining -= len(selected)
        complete = len(selected) == len(population)
        current = {'name': sheet.title, 'visibility': sheet.sheet_state, 'dimensions': sheet.calculate_dimension(),
                   'nonEmptyCells': len(population), 'profiledCells': len(selected), 'complete': complete,
                   'mergedRanges': [str(r) for r in sheet.merged_cells.ranges], 'mergedRangesTotal': len(sheet.merged_cells.ranges),
                   'hiddenRows': [row for row, dim in sheet.row_dimensions.items() if dim.hidden],
                   'hiddenColumns': [column for column, dim in sheet.column_dimensions.items() if dim.hidden],
                   'formulas': sum(cell.data_type == 'f' for cell in population), 'comments': sum(bool(cell.comment) for cell in population),
                   'images': len(sheet._images), 'styles': sorted(set(cell.style_id for cell in population)), 'regions': [], 'regionsOmitted': 0}
        if not complete:
            result['complete'] = False
            result['warnings'].append('Sheet profile limit: ' + sheet.title)
        by_row = {}
        for cell in selected:
            by_row.setdefault(cell.row, []).append(cell)
        for row_band in bands(by_row):
            row_cells = [cell for row in row_band for cell in by_row[row]]
            by_column = {}
            for cell in row_cells:
                by_column.setdefault(cell.column, []).append(cell)
            for column_band in bands(by_column):
                block = sorted((cell for column in column_band for cell in by_column[column]), key=lambda cell: (cell.row, cell.column))
                if region_count >= 300:
                    current['regionsOmitted'] += 1
                    current['complete'] = False
                    result['complete'] = False
                    continue
                minimum_column = min(cell.column for cell in block)
                maximum_column = max(cell.column for cell in block)
                minimum_row = min(cell.row for cell in block)
                maximum_row = max(cell.row for cell in block)
                count = min(24, len(block))
                indices = sorted(set(round(i * (len(block)-1) / max(1, count-1)) for i in range(count)))
                current['regions'].append({'id': 's%d-r%d' % (index+1, len(current['regions'])+1),
                    'range': '%s%d:%s%d' % (get_column_letter(minimum_column), minimum_row, get_column_letter(maximum_column), maximum_row),
                    'nonEmptyCells': len(block), 'sample': [record(block[i], cached[sheet.title], sheet) for i in indices], 'sampleOnly': len(block) > count})
                region_count += 1
        if current['regionsOmitted']:
            result['warnings'].append('Region profile limit: ' + sheet.title)
        result['sheets'].append(current)
    return result

def strict_object(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            raise ValueError('duplicate key')
        result[key] = value
    return result

def details(workbook, cached, request):
    if set(request) != {'ranges'} or not isinstance(request['ranges'], list) or not 1 <= len(request['ranges']) <= 18:
        raise ValueError('invalid ranges')
    result = {'ranges': [], 'complete': True, 'cellsRead': 0}
    seen = set()
    characters = 0
    for selection in request['ranges']:
        if set(selection) != {'sheet', 'range'} or selection['sheet'] not in workbook.sheetnames:
            raise ValueError('invalid range keys')
        bounds = range_boundaries(selection['range'].replace('$', '').upper())
        left, top, right, bottom = bounds
        if None in bounds or not 1 <= left <= right <= 16384 or not 1 <= top <= bottom <= 1048576 or (right-left+1)*(bottom-top+1) > 12000:
            raise ValueError('range bounds')
        sheet = workbook[selection['sheet']]
        output = {**selection, 'cells': [], 'truncated': False}
        for cell in stored_cells(sheet):
            if not left <= cell.column <= right or not top <= cell.row <= bottom:
                continue
            value = record(cell, cached[sheet.title], sheet)
            key = (sheet.title, cell.coordinate)
            cost = len(json.dumps(value, ensure_ascii=False))
            if key not in seen:
                if len(seen) >= 6000 or characters + cost > 220000:
                    output['truncated'] = True
                    result['complete'] = False
                    break
                seen.add(key)
                characters += cost
            output['cells'].append(value)
            if value.get('truncated'):
                output['truncated'] = True
                result['complete'] = False
        result['ranges'].append(output)
    result['cellsRead'] = len(seen)
    return result

def main():
    source_path, operation, output_path = sys.argv[1:4]
    with open(source_path, 'rb') as handle:
        data = handle.read(20*1024*1024+1)
    if not data or len(data) > 20*1024*1024:
        raise ValueError('source size')
    archive = zipfile.ZipFile(io.BytesIO(data))
    if len(archive.infolist()) > 20000 or sum(e.file_size for e in archive.infolist()) > 100*1024*1024:
        raise ValueError('archive limit')
    workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=False, read_only=False, keep_links=False)
    cached = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=False, keep_links=False)
    if operation == 'inventory':
        result = inventory(workbook, cached)
    elif operation == 'read':
        with open(sys.argv[4], encoding='utf-8') as handle:
            request = json.load(handle, object_pairs_hook=strict_object)
        result = details(workbook, cached, request)
    else:
        raise ValueError('operation')
    with open(output_path, 'w', encoding='utf-8') as handle:
        json.dump(result, handle, ensure_ascii=False)
    workbook.close()
    cached.close()
    print(json.dumps({'cellsRead': result.get('cellsRead', sum(s['profiledCells'] for s in result.get('sheets', [])))}))

if __name__ == '__main__':
    main()

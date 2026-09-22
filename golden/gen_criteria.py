#!/usr/bin/env python
"""기준서(criteria) 골든 픽스처 생성기.

golden/criteria/C01.xlsx ... C20.xlsx 와 동일 stem 의 .json (ground truth) 을 만든다.
실행: uv run python golden/gen_criteria.py   (멱등 - 항상 덮어쓴다)
"""
from __future__ import annotations

import io
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font
from PIL import Image as PILImage, ImageDraw, ImageFont

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

OUT = Path(__file__).resolve().parent / "criteria"
BOLD = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center")


# --------------------------------------------------------------------------- #
# criteria pool
# --------------------------------------------------------------------------- #
def C(item, raw, op, value, unit, note="", method="KS L 5405", unit_cell=True):
    """하나의 기준 행. raw = 셀에 찍히는 문자열, op/value/unit = 정답."""
    return {"item": item, "raw": raw, "op": op, "value": value, "unit": unit,
            "note": note, "method": method, "unit_cell": unit_cell}


BASE = [
    C("산화마그네슘", "2.0 이상", ">=", 2.0, "%", "MgO"),
    C("염화물", "0.02 이하", "<=", 0.02, "%", "Cl-"),
    C("강열감량", "5.0 이하", "<=", 5.0, "%"),
    C("비표면적", "4000 이상", ">=", 4000, "cm2/g", method="KS F 2563"),
    C("활성도지수", "80 이상", ">=", 80, "%", "28일", method="KS F 2563"),
    C("밀도", "2.20 이상", ">=", 2.20, "g/cm3", method="KS F 2563"),
    C("압축강도", "24.0 이상", ">=", 24.0, "MPa", "28일", method="KS F 2405"),
]

CHEM, PHYS = BASE[:3], BASE[3:]

# rules.json 의 20가지 기준 표현형을 모두 담은 목록 (C10, C19 에서 사용)
MIXED = [
    C("산화마그네슘", "2.0 이상", ">=", 2.0, "%"),
    C("염화물", "0.02 이하", "<=", 0.02, "%"),
    C("알칼리량", "1.0 초과", ">", 1.0, "%"),
    C("강열감량", "5.0 미만", "<", 5.0, "%"),
    C("수분함량", "2~5", "between", [2, 5], "%"),
    C("삼산화황", "0.5 이상 1.5 이하", "between", [0.5, 1.5], "%"),
    C("유리석회", "\u22640.02", "<=", 0.02, "%"),
    C("밀도", "\u22652.20", ">=", 2.20, "g/cm3"),
    C("팽창률", "<0.10", "<", 0.10, "%"),
    C("활성도지수", ">80", ">", 80, "%"),
    C("비표면적", "min 4000", ">=", 4000, "cm2/g"),
    C("응결시간(초결)", "max 300", "<=", 300, "min"),
    C("공기량", "3.0% 이상", ">=", 3.0, "%", unit_cell=False),
    C("압축강도", "24.0 이상(3회 평균)", ">=", 24.0, "MPa"),
    C("유해물질", "N.D.", "qualitative", None, None),
    C("6가크롬", "불검출", "qualitative", None, None),
    C("외관", "이상 없을 것", "qualitative", None, None),
    C("안정성", "적합할 것", "qualitative", None, None),
    C("부가시험", "-", "qualitative", None, None),
    C("규격적합성", "KS F 2563 만족", "qualitative", None, None),
]


# --------------------------------------------------------------------------- #
# sheet helpers
# --------------------------------------------------------------------------- #
def head(ws, row, cols, labels):
    for col, lab in zip(cols, labels):
        cell = ws[f"{col}{row}"]
        cell.value = lab
        cell.font = BOLD
        cell.alignment = CENTER


def rows(ws, start, cols, items, method_col=None):
    """cols = (item, op_value, unit, note); unit/note 는 None 이면 생략."""
    ic, vc, uc, nc = cols
    r = start
    for it in items:
        ws[f"{ic}{r}"] = it["item"]
        ws[f"{vc}{r}"] = it["raw"]
        if uc and it["unit"] and it["unit_cell"]:
            ws[f"{uc}{r}"] = it["unit"]
        if nc and it["note"]:
            ws[f"{nc}{r}"] = it["note"]
        if method_col:
            ws[f"{method_col}{r}"] = it["method"]
        r += 1
    return r


def kfont(size):
    for p in (r"C:\Windows\Fonts\malgun.ttf",
              "/usr/share/fonts/truetype/nanum/NanumGothic.ttf",
              "/System/Library/Fonts/Supplemental/AppleGothic.ttf"):
        try:
            if Path(p).exists():
                return ImageFont.truetype(p, size)
        except Exception:
            pass
    return ImageFont.load_default()


def logo_png(w=96, h=32):
    img = PILImage.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    d.rectangle([0, 0, w - 1, h - 1], outline=(20, 70, 140), width=2)
    d.text((10, 8), "GS-QM", fill=(20, 70, 140), font=kfont(14))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return buf


def criteria_png(items, w=560, h=250):
    """기준표 자체를 그림으로 렌더 (C12 용)."""
    img = PILImage.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    f, fb = kfont(14), kfont(15)
    d.text((12, 8), "품질기준표 (스캔본)", fill="black", font=fb)
    xs = [12, 210, 360, 460]
    y = 40
    for x, lab in zip(xs, ("시험항목", "기준", "단위", "비고")):
        d.text((x, y), lab, fill="black", font=fb)
    d.line([(8, y + 20), (w - 8, y + 20)], fill="black", width=1)
    y += 28
    for it in items:
        d.text((xs[0], y), it["item"], fill="black", font=f)
        d.text((xs[1], y), it["raw"], fill="black", font=f)
        d.text((xs[2], y), it["unit"] or "-", fill="black", font=f)
        y += 24
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return buf


# --------------------------------------------------------------------------- #
# ground truth builders
# --------------------------------------------------------------------------- #
def gt_adapter(cid, desc, sheet, header_row, data_from, cols, items, **extra):
    return {
        "id": cid, "desc": desc, "expect": "adapter", "sheet": sheet,
        "table": {"header_row": header_row, "data_from": data_from},
        "columns": {"item": cols[0], "op_value": cols[1],
                    "unit": cols[2], "note": cols[3]},
        "criteria": [{"item": i["item"], "op": i["op"],
                      "value": i["value"], "unit": i["unit"]} for i in items],
        **extra,
    }


def gt_flag(cid, desc, sheet, reason, **extra):
    return {"id": cid, "desc": desc, "expect": "flag", "sheet": sheet,
            "flag_reason": reason, **extra}


def save(wb, gt):
    OUT.mkdir(parents=True, exist_ok=True)
    xp = OUT / f"{gt['id']}.xlsx"
    jp = OUT / f"{gt['id']}.json"
    wb.save(xp)
    jp.write_text(json.dumps(gt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    hr = gt.get("table", {}).get("header_row", "-")
    print(f"[criteria] {xp.name}  expect={gt['expect']}  sheet={gt['sheet']}  "
          f"header_row={hr}  criteria={len(gt.get('criteria', []))}  :: {gt['desc']}")


# --------------------------------------------------------------------------- #
# C01 ~ C20
# --------------------------------------------------------------------------- #
def c01():
    wb = Workbook()
    ws = wb.active
    ws.title = "품질기준"
    cols = ("A", "B", "C", "D")
    head(ws, 1, cols, ("시험항목", "기준", "단위", "비고"))
    rows(ws, 2, cols, BASE)
    return wb, gt_adapter("C01", "A1에서 시작하는 깨끗한 표, 헤더 1행",
                          "품질기준", 1, 2, cols, BASE)


def c02():
    wb = Workbook()
    ws = wb.active
    ws.title = "품질기준"
    ws["A1"] = "[LOGO]"
    ws.merge_cells("A2:F2")
    ws["A2"] = "고로슬래그 미분말 품질기준서 (2026년 개정)"
    ws["A2"].font = Font(bold=True, size=14)
    ws["A2"].alignment = CENTER
    ws["A4"] = "※ 본 기준은 KS F 2563 을 근거로 하며, 현장 반입 자재에 한하여 적용한다."
    cols = ("B", "E", "D", "F")
    head(ws, 8, ("B", "C", "D", "E", "F"),
         ("시험항목", "시험방법", "단위", "기준", "비고"))
    rows(ws, 9, cols, BASE, method_col="C")
    return wb, gt_adapter("C02", "로고·병합제목·안내문·빈행 뒤 8행 헤더",
                          "품질기준", 8, 9, cols, BASE)


def c03():
    wb = Workbook()
    ws = wb.active
    ws.title = "품질기준"
    ws["C2"] = "○○현장 자재 품질관리 기준"
    ws["C4"] = "문서번호 : QM-2026-014"
    ws["C5"] = "개정일자 : 2026-01-05"
    ws["C7"] = "1. 적용범위"
    ws["C8"] = "  본 기준은 콘크리트용 혼화재료에 적용한다."
    ws["C10"] = "2. 품질기준"
    cols = ("C", "D", "E", "F")
    head(ws, 12, cols, ("시험항목", "기준", "단위", "비고"))
    rows(ws, 13, cols, BASE)
    return wb, gt_adapter("C03", "헤더 12행 + 표가 C열에서 시작 (A,B열 비어있음)",
                          "품질기준", 12, 13, cols, BASE)


def c04():
    wb = Workbook()
    ws = wb.active
    ws.title = "품질기준"
    ws["B3"] = "품질기준표"
    ws["B3"].font = Font(bold=True, size=13)
    ws["B7"] = "시험항목"
    ws["C7"] = "품질기준"
    ws["E7"] = "비고"
    ws["C8"] = "기준"
    ws["D8"] = "단위"
    for ref in ("B7:B8", "C7:D7", "E7:E8"):
        ws.merge_cells(ref)
    for ref in ("B7", "C7", "E7", "C8", "D8"):
        ws[ref].font = BOLD
        ws[ref].alignment = CENTER
    cols = ("B", "C", "D", "E")
    rows(ws, 9, cols, BASE)
    return wb, gt_adapter("C04", "2행 병합 헤더 (7행 그룹헤더 / 8행 소헤더)",
                          "품질기준", 8, 9, cols, BASE)


def c05():
    wb = Workbook()
    cover = wb.active
    cover.title = "표지"
    cover["B1"] = "자재 품질관리 기준서"
    cover["B1"].font = Font(bold=True, size=16)
    head(cover, 3, ("B", "C", "D", "E"), ("시험항목", "기준", "단위", "비고"))
    for i, (a, b) in enumerate([("문서번호", "QM-2026-014"), ("작성일자", "2026-01-05"),
                                ("작성자", "김○○"), ("승인자", "박○○")]):
        cover[f"B{4 + i}"] = a
        cover[f"C{4 + i}"] = b
        cover[f"D{4 + i}"] = "-"

    ws = wb.create_sheet("품질기준")
    cols = ("A", "B", "C", "D")
    head(ws, 1, cols, ("시험항목", "기준", "단위", "비고"))
    rows(ws, 2, cols, BASE)

    hist = wb.create_sheet("변경이력")
    head(hist, 1, ("A", "B", "C"), ("개정번호", "개정일자", "개정내용"))
    for i, (a, b, c) in enumerate([("0", "2025-03-11", "최초 제정"),
                                   ("1", "2026-01-05", "활성도지수 기준 상향")]):
        hist[f"A{2 + i}"] = a
        hist[f"B{2 + i}"] = b
        hist[f"C{2 + i}"] = c
    return wb, gt_adapter("C05", "3시트(표지 미끼표 / 품질기준 실제 / 변경이력)",
                          "품질기준", 1, 2, cols, BASE)


def c06():
    wb = Workbook()
    ws = wb.active
    ws.title = "품질기준"
    ws["A1"] = "고로슬래그 미분말 품질기준"
    ws["A1"].font = Font(bold=True, size=13)
    cols = ("A", "B", "C", "D")
    head(ws, 5, cols, ("시험항목", "기준", "단위", "비고"))
    ws["A6"] = "[화학성분]"
    ws["A6"].font = BOLD
    rows(ws, 7, cols, CHEM)          # 7,8,9 / 10행 비움
    ws["A11"] = "[물리성능]"
    ws["A11"].font = BOLD
    rows(ws, 12, cols, PHYS)         # 12..15
    return wb, gt_adapter("C06", "구역 캡션 [화학성분]/[물리성능] + 사이 빈행, 헤더는 5행 하나",
                          "품질기준", 5, 6, cols, BASE)


def c07():
    wb = Workbook()
    ws = wb.active
    ws.title = "품질기준"
    ws["A1"] = "품질기준 및 검사항목"
    ws["A1"].font = Font(bold=True, size=13)
    cols = ("A", "B", "C", "D")
    head(ws, 3, cols, ("시험항목", "기준", "단위", "비고"))
    r = rows(ws, 4, cols, CHEM)
    ws[f"A{r}"] = "소계"
    ws[f"B{r}"] = "3개 항목"
    ws[f"A{r}"].font = BOLD
    r = rows(ws, r + 1, cols, PHYS)
    ws[f"A{r}"] = "합계"
    ws[f"B{r}"] = "7개 항목"
    ws[f"A{r}"].font = BOLD
    return wb, gt_adapter("C07", "데이터 중간·끝에 소계/합계 행 (중지규칙에서 제외되어야 함)",
                          "품질기준", 3, 4, cols, BASE)


def c08():
    wb = Workbook()
    ws = wb.active
    ws.title = "품질기준"
    ws["A1"] = "품질기준"
    ws["A1"].font = Font(bold=True, size=13)
    cols = ("A", "B", "C", "D")
    head(ws, 2, cols, ("시험항목", "기준", "단위", "비고"))
    end = rows(ws, 3, cols, BASE)          # data 3..9, end == 10
    ws["B5"].comment = Comment("2026년 개정 시 수치 재확인 요망", "김○○")
    ws[f"F{end + 2}"] = "← 작성: 김○○ (2026-01-05)"
    return wb, gt_adapter("C08", "데이터 3행 아래 모서리 메모 셀 + 데이터 셀에 실제 코멘트",
                          "품질기준", 2, 3, cols, BASE)


def c09():
    wb = Workbook()
    ws = wb.active
    ws.title = "품질기준"
    items = [
        C("산화마그네슘", "2.0 이상", ">=", 2.0, "%"),                      # 단위 열
        C("염화물", "0.02% 이하", "<=", 0.02, "%", unit_cell=False),         # 값 안의 단위
        C("강열감량", "5.0 이하", "<=", 5.0, "%", unit_cell=False),          # 헤더의 단위
    ]
    cols = ("A", "B", "C", "D")
    head(ws, 1, cols, ("시험항목", "기준(%)", "단위", "비고"))
    rows(ws, 2, cols, items)
    return wb, gt_adapter("C09", "단위가 3곳(단위열 / 값 안 / 헤더)에 흩어져 있음",
                          "품질기준", 1, 2, cols, items)


def c10():
    wb = Workbook()
    ws = wb.active
    ws.title = "품질기준"
    cols = ("A", "B", "C", "D")
    head(ws, 1, cols, ("시험항목", "기준", "단위", "비고"))
    rows(ws, 2, cols, MIXED)
    return wb, gt_adapter("C10", "rules.json 의 20가지 기준 표현형이 모두 섞임",
                          "품질기준", 1, 2, cols, MIXED)


def c11():
    wb = Workbook()
    ws = wb.active
    ws.title = "품질기준"
    img = XLImage(logo_png())
    img.anchor = "A1"
    ws.add_image(img)
    ws["C2"] = "고로슬래그 미분말 품질기준"
    ws["C2"].font = Font(bold=True, size=13)
    cols = ("A", "B", "C", "D")
    head(ws, 6, cols, ("시험항목", "기준", "단위", "비고"))
    rows(ws, 7, cols, BASE)
    return wb, gt_adapter("C11", "A1에 로고 이미지 삽입 + 아래에 정상 표",
                          "품질기준", 6, 7, cols, BASE)


def c12():
    wb = Workbook()
    ws = wb.active
    ws.title = "품질기준"
    ws["A1"] = "품질기준 (스캔 첨부)"
    ws["A1"].font = Font(bold=True, size=13)
    img = XLImage(criteria_png(BASE))
    img.anchor = "A3"
    ws.add_image(img)
    return wb, gt_flag("C12", "기준이 삽입 이미지 안에만 존재, 셀은 비어있음",
                       "품질기준", "기준 표가 이미지로만 존재하여 셀에서 읽을 수 없음")


def c13():
    wb = Workbook()
    ws = wb.active
    ws.title = "옵션표"
    ws["A1"] = "○○아파트 101동 타입별 마감 옵션표"
    ws["A1"].font = Font(bold=True, size=13)
    head(ws, 3, ("B", "C", "D", "E", "F"), ("층", "101호", "102호", "103호", "104호"))
    codes = ["A-1", "A-2", "B-1", "B-2", "C-1"]
    for i, floor in enumerate(range(15, 0, -1)):
        r = 4 + i
        ws[f"B{r}"] = floor
        for j, col in enumerate("CDEF"):
            ws[f"{col}{r}"] = codes[(i + j) % len(codes)]
    return wb, gt_flag("C13", "동/호 그리드(층 15..1 × 101~104호) 배치표",
                       "옵션표", "아파트 동/호 배치표이며 품질기준 표가 아님")


def c14():
    wb = Workbook()
    ws = wb.active
    ws.title = "품질기준"
    cols = ("A", "B", "C", "D")
    head(ws, 1, cols, ("시험항목", "기준", "단위", "비고"))
    rows(ws, 2, cols, BASE)
    # 기준 값은 리터럴 유지. 수식은 비고(D) 열에만 넣는다.
    # (openpyxl 은 캐시값을 쓸 수 없으므로 data_only 로 읽으면 D열은 None 이 된다.)
    ws["D2"] = "=ROUND(2.0,1)"
    ws["D3"] = "=ROUND(0.02,3)"
    ws["D4"] = '=A4&" 확인"'
    ws["D5"] = "=4000*1"
    return wb, gt_adapter("C14", "비고 열에만 수식(=ROUND 등), 기준 값은 리터럴 유지",
                          "품질기준", 1, 2, cols, BASE)


def c15():
    wb = Workbook()
    ws = wb.active
    ws.title = "품질기준"
    cols = ("A", "B", "C", "E")
    head(ws, 1, ("A", "B", "C", "D", "E"),
         ("시험항목", "기준", "단위", "시험방법", "비고"))
    rows(ws, 2, cols, BASE, method_col="D")
    for r in (4, 6):
        ws.row_dimensions[r].hidden = True
    ws.column_dimensions["D"].hidden = True
    return wb, gt_adapter("C15", "데이터 중간 2행 숨김(4,6행) + 시험방법 D열 숨김",
                          "품질기준", 1, 2, cols, BASE)


def c16():
    wb = Workbook()
    ws = wb.active
    ws.title = "품질기준"
    items = [BASE[0], C("산화마그네슘", "5.0 이하", "<=", 5.0, "%", "상한")] + BASE[1:]
    cols = ("A", "B", "C", "D")
    head(ws, 1, cols, ("시험항목", "기준", "단위", "비고"))
    rows(ws, 2, cols, items)
    return wb, gt_adapter("C16", "동일 항목명이 서로 다른 한계값으로 2회 등장",
                          "품질기준", 1, 2, cols, items,
                          uncertain=["duplicate item: 산화마그네슘"])


def c17():
    wb = Workbook()
    ws = wb.active
    ws.title = "품질기준"
    cols = ("A", "B", "C", "D")
    head(ws, 1, cols, ("  시험항목 ", "기\u200b준", " 단위\u200b", "비 고 "))
    rows(ws, 2, cols, BASE)
    return wb, gt_adapter("C17", "헤더 라벨에 공백·zero-width(U+200B) 문자 삽입",
                          "품질기준", 1, 2, cols, BASE)


def c18():
    wb = Workbook()
    ws = wb.active
    ws.title = "품질기준"
    ws["A1"] = "2026년 공정 계획표 (40열)"
    ws["A1"].font = Font(bold=True, size=13)
    for i in range(1, 41):
        ws.cell(row=2, column=i, value=f"W{i}")
        ws.cell(row=3, column=i, value=(i * 3) % 17)
    cols = ("AA", "AB", "AC", "AE")
    head(ws, 6, ("AA", "AB", "AC", "AD", "AE"),
         ("시험항목", "기준", "단위", "시험방법", "비고"))
    rows(ws, 7, cols, BASE, method_col="AD")
    return wb, gt_adapter("C18", "시트 폭 40열, 기준표는 AA~AE 열에 위치",
                          "품질기준", 6, 7, cols, BASE)


def c19():
    wb = Workbook()
    ws = wb.active
    ws.title = "품질기준"
    cols = ("A", "B", "C", "D")
    head(ws, 1, cols, ("시험항목", "기준", "단위", "비고"))
    rows(ws, 2, cols, MIXED)                     # 2..21
    start = 23                                   # 22행은 빈 행
    for i in range(4980):
        r = start + i
        ws.cell(row=r, column=1, value=f"2026-{(i % 12) + 1:02d}-{(i % 28) + 1:02d}")
        ws.cell(row=r, column=2, value=f"설비{(i % 7) + 1}")
        ws.cell(row=r, column=3, value=round(18.0 + (i % 130) * 0.1, 1))
        ws.cell(row=r, column=4, value="정상" if i % 53 else "점검")
    return wb, gt_adapter("C19", "20개 기준행 뒤 빈 행, 그 아래 4980행의 무관한 로그 데이터",
                          "품질기준", 1, 2, cols, MIXED)


def c20():
    wb = Workbook()
    ws = wb.active
    ws.title = "품질기준"
    ws["A1"] = "품질기준 (코드 기준)"
    ws["A1"].font = Font(bold=True, size=13)
    cols = ("A", "B", "C", "D")
    head(ws, 3, cols, ("코드", "기준", "단위", "비고"))
    for i, it in enumerate(BASE):
        r = 4 + i
        ws[f"A{r}"] = f"Q-{i + 1:02d}"
        ws[f"B{r}"] = it["raw"]
        if it["unit"]:
            ws[f"C{r}"] = it["unit"]
        if it["note"]:
            ws[f"D{r}"] = it["note"]
    code = wb.create_sheet("코드표")
    head(code, 1, ("A", "B"), ("코드", "시험항목"))
    for i, it in enumerate(BASE):
        code[f"A{2 + i}"] = f"Q-{i + 1:02d}"
        code[f"B{2 + i}"] = it["item"]
    return wb, gt_adapter("C20", "품질기준 시트는 코드만, 항목명은 코드표 시트에 존재",
                          "품질기준", 3, 4, cols, BASE,
                          uncertain=["item names live in 코드표 sheet"],
                          code_map={"sheet": "코드표", "code_column": "A",
                                    "name_column": "B", "header_row": 1})


BUILDERS = [c01, c02, c03, c04, c05, c06, c07, c08, c09, c10,
            c11, c12, c13, c14, c15, c16, c17, c18, c19, c20]


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    for fn in BUILDERS:
        wb, gt = fn()
        save(wb, gt)
        wb.close()
    print(f"[criteria] done - {len(BUILDERS)} xlsx + {len(BUILDERS)} json in {OUT}")


if __name__ == "__main__":
    main()

#!/usr/bin/env python
"""검토대장(ledger) 골든 픽스처 생성기.

golden/ledger/L01.xlsx ... L06.xlsx 와 동일 stem 의 .json (ground truth) 을 만든다.
실행: uv run python golden/gen_ledger.py   (멱등 - 항상 덮어쓴다)
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

OUT = Path(__file__).resolve().parent / "ledger"
SHEET = "검토대장"
KEY = "2026-0891"
BOLD = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center")

MATERIALS = ["고로슬래그 미분말", "플라이애시", "1종 보통포틀랜드시멘트", "레디믹스트 콘크리트",
             "철근 SD400", "고로슬래그 미분말", "혼화제(AE감수제)", "잔골재",
             "굵은골재", "플라이애시", "고로슬래그 미분말", "1종 보통포틀랜드시멘트"]
N_ROWS = 12
KEY_INDEX = 10          # 0-based -> header_row 1 기준 12행


def make_rows():
    """(no, 반입일, 성적서번호, 자재, 수량, 판정, 비고) 12행."""
    out = []
    for i in range(N_ROWS):
        cert = f"2026-{881 + i:04d}"
        pending = i == KEY_INDEX
        out.append((
            i + 1,
            f"2026-09-{i + 1:02d}",
            cert,
            MATERIALS[i],
            (i + 1) * 20,
            "" if pending else "적합",
            "" if pending else "이상없음",
        ))
    return out


def head(ws, row, cols, labels):
    for col, lab in zip(cols, labels):
        cell = ws[f"{col}{row}"]
        cell.value = lab
        cell.font = BOLD
        cell.alignment = CENTER


def gt(lid, desc, expect, cols, key_header, result_col, note_col,
       target_cells=None, **extra):
    d = {"id": lid, "desc": desc, "expect": expect, "sheet": SHEET, "header_row": 1,
         "key_column": cols, "key_header": key_header,
         "result_column": result_col, "note_column": note_col, "key": KEY}
    if target_cells is not None:
        d["target_cells"] = target_cells
    d.update(extra)
    return d


def save(wb, g):
    OUT.mkdir(parents=True, exist_ok=True)
    xp = OUT / f"{g['id']}.xlsx"
    jp = OUT / f"{g['id']}.json"
    wb.save(xp)
    jp.write_text(json.dumps(g, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    tgt = ",".join(g.get("target_cells", [])) or "-"
    print(f"[ledger]   {xp.name}  expect={g['expect']}  key_col={g['key_column']}  "
          f"result_col={g['result_column']}  target={tgt}  :: {g['desc']}")


# --------------------------------------------------------------------------- #
# L01 ~ L06
# --------------------------------------------------------------------------- #
def plain_sheet(wb=None):
    """A No. / B 반입일 / C 성적서번호 / D 자재 / E 수량 / F 판정 / G 비고."""
    wb = wb or Workbook()
    ws = wb.active
    ws.title = SHEET
    cols = ("A", "B", "C", "D", "E", "F", "G")
    head(ws, 1, cols, ("No.", "반입일", "성적서번호", "자재", "수량", "판정", "비고"))
    for i, row in enumerate(make_rows()):
        for col, val in zip(cols, row):
            if val != "":
                ws[f"{col}{2 + i}"] = val
    return wb, ws


def l01():
    wb, _ = plain_sheet()
    return wb, gt("L01", "기본", "ok", "C", "성적서번호", "F", "G", ["F12", "G12"])


def l02():
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    cols = ("A", "B", "C", "D", "E", "F", "G", "H", "I")
    head(ws, 1, cols, ("No.", "반입일", "공정", "자재", "성적서번호",
                       "규격", "수량", "판정", "비고"))
    for i, (no, day, cert, mat, qty, res, note) in enumerate(make_rows()):
        r = 2 + i
        ws[f"A{r}"] = no
        ws[f"B{r}"] = day
        ws[f"C{r}"] = "골조공사"
        ws[f"D{r}"] = mat
        ws[f"E{r}"] = cert
        ws[f"F{r}"] = "KS F 2563"
        ws[f"G{r}"] = qty
        if res:
            ws[f"H{r}"] = res
        if note:
            ws[f"I{r}"] = note
    return wb, gt("L02", "키 열 E, 판정 열 H", "ok", "E", "성적서번호", "H", "I",
                  ["H12", "I12"])


def l03():
    wb, ws = plain_sheet()
    ws["C5"] = KEY                     # 같은 키가 5행에도 등장
    ws["F5"] = ""
    ws["G5"] = ""
    return wb, gt("L03", "동일 성적서번호가 2개 행(5행, 12행)에 존재", "flag",
                  "C", "성적서번호", "F", "G", None,
                  flag_reason="duplicate key", duplicate_rows=[5, 12])


def l04():
    wb, ws = plain_sheet()
    ws["A15"] = "합계"
    ws["A15"].font = BOLD
    ws["E15"] = "=SUM(E2:E13)"
    return wb, gt("L04", "데이터 아래 =SUM() 합계 행 (건드리면 안 됨)", "ok",
                  "C", "성적서번호", "F", "G", ["F12", "G12"])


def l05():
    wb, ws = plain_sheet()
    ws["B15"] = "← 09.03 확인 (박○○)"
    ws["A17"] = "합계"
    ws["A17"].font = BOLD
    ws["E17"] = "=SUM(E2:E13)"
    return wb, gt("L05", "데이터와 합계 사이에 메모 행", "ok",
                  "C", "성적서번호", "F", "G", ["F12", "G12"])


def l06():
    wb, ws = plain_sheet()
    ws.protection.sheet = True
    return wb, gt("L06", "시트 보호 설정됨", "flag",
                  "C", "성적서번호", "F", "G", ["F12", "G12"],
                  flag_reason="protected sheet")


BUILDERS = [l01, l02, l03, l04, l05, l06]


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    for fn in BUILDERS:
        wb, g = fn()
        save(wb, g)
        wb.close()
    print(f"[ledger]   done - {len(BUILDERS)} xlsx + {len(BUILDERS)} json in {OUT}")


if __name__ == "__main__":
    main()

"""Physical openpyxl inputs for the actual criteria profile module; no provider/oracle access."""
import gc
import hashlib
import importlib.util
import json
from pathlib import Path
import sys

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill


def main():
    root, destination = map(Path, sys.argv[1:3])
    destination.mkdir(parents=True, exist_ok=False)
    product_path = root / 'server' / 'criteria-workbook-profile.py'
    spec = importlib.util.spec_from_file_location('actual_criteria_workbook_profile', product_path)
    product = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(product)
    rows = []

    def observe(identifier, workbook, operation, request=None, recipe=None, existing_input=None):
        filename = Path(existing_input) if existing_input else destination / (identifier + '.xlsx')
        if existing_input is None:
            workbook.save(filename)
        # Reopen the real artifact in both formula and cached modes, exactly as production main does.
        actual = openpyxl.load_workbook(filename, data_only=False, keep_links=False)
        cached = openpyxl.load_workbook(filename, data_only=True, keep_links=False)
        try:
            try:
                output = product.inventory(actual, cached) if operation == 'inventory' else product.details(actual, cached, request)
                result = {'accepted': True, 'output': output}
            except Exception as error:
                result = {'accepted': False, 'error': {'type': type(error).__name__, 'message': str(error)}}
            row = {'id': identifier, 'input': {'recipe': recipe, 'operation': operation, 'request': request,
                   'workbook': str(filename.relative_to(root)).replace('\\', '/'),
                   'sha256': hashlib.sha256(filename.read_bytes()).hexdigest()}, 'result': result}
            rows.append(row)
            return row
        finally:
            actual.close()
            cached.close()
            # Workbook/cell parent cycles must be unreachable before collection,
            # otherwise the next 200k-cell boundary would retain both prior copies.
            actual = None
            cached = None
            gc.collect()

    def new():
        value = openpyxl.Workbook()
        value.active.title = 'Rules'
        return value

    for count in (99, 100, 101):
        workbook = new()
        for number in range(1, count):
            workbook.create_sheet('Sheet-' + str(number + 1))
        for sheet in workbook:
            sheet['A1'] = 'source'
        observe('sheets-' + str(count), workbook, 'inventory', recipe={'sheets': count, 'populatedPerSheet': 1})
        workbook.close()
    workbook = new()
    for number in range(1, 103):
        workbook.create_sheet('Sheet-' + str(number + 1))
    for sheet in workbook.worksheets[:100]:
        sheet['A1'] = 'source'
    workbook.worksheets[101]['A1'] = 'beyond profile'
    workbook.worksheets[102]['A1'].comment = Comment('comment-only beyond profile', 'Synthetic')
    observe('sheets-tail', workbook, 'inventory', recipe={'sheets': 103, 'tail': ['empty', 'populated', 'comment-only']})
    workbook.close()

    # Genuine populated cells; no monkeypatch of production caps or record serialization.
    workbook = new()
    for number in range(200001):
        workbook.active.cell(number // 1000 + 1, number % 1000 + 1, 'x')
    observe('cells-200001', workbook, 'inventory', recipe={'populatedCells': 200001, 'columns': 1000, 'value': 'x'})
    workbook.active.cell(201, 1).value = None
    observe('cells-200000', workbook, 'inventory', recipe={'populatedCells': 200000, 'columns': 1000, 'value': 'x'})
    workbook.close()
    del workbook
    gc.collect()

    for count in (300, 301):
        workbook = new()
        for number in range(count):
            workbook.active.cell(number * 2 + 1, 1, 'x')
        observe('regions-' + str(count), workbook, 'inventory', recipe={'isolatedRows': count, 'gap': 1})
        workbook.close()

    workbook = new()
    for number in range(49):
        workbook.active.cell(number + 1, 1, 'sample-' + str(number + 1))
    observe('sample-49', workbook, 'inventory', recipe={'populatedCells': 49, 'column': 'A', 'sourceOrder': 'ascending row'})
    workbook.close()

    for cost in (219999, 220000, 220001):
        workbook = new()
        sheet = workbook.active
        for number in range(1, 61):
            sheet.cell(number, 1, '')
        overhead = sum(len(json.dumps(product.record(cell, sheet, sheet), ensure_ascii=False)) for cell in sheet['A'])
        remaining = cost - overhead
        lengths = []
        for number in range(1, 61):
            # Every cell stays populated and below the per-cell 4,000-character limit.
            length = min(3900, remaining - (60 - number))
            sheet.cell(number, 1, '한' * length)
            remaining -= length
            lengths.append(length)
        assert remaining == 0 and all(1 <= length <= 4000 for length in lengths)
        exact = sum(len(json.dumps(product.record(cell, sheet, sheet), ensure_ascii=False)) for cell in sheet['A'])
        assert exact == cost
        observe('characters-' + str(cost), workbook, 'read', {'ranges': [{'sheet': 'Rules', 'range': 'A1:A60'}]},
                {'recordJsonCharacters': cost, 'textLengths': lengths, 'character': '한', 'ensureAscii': False})
        workbook.close()

    workbook = new()
    for number in range(1, 6002):
        workbook.active.cell(number, 1, 'x')
    row = observe('cells-6001-detail', workbook, 'read', {'ranges': [{'sheet': 'Rules', 'range': 'A1:A6001'}]},
                  {'populatedCells': 6001, 'value': 'x', 'unchangedProductionSerialization': True})
    raw = row['result']['output']
    costs = [len(json.dumps(product.record(cell, workbook.active, workbook.active), ensure_ascii=False)) for cell in workbook.active['A']]
    row['independentRecordAccounting'] = {'all6001Characters': sum(costs), 'retainedCharacters': sum(costs[:raw['cellsRead']]),
                                        'plusNextCharacters': sum(costs[:raw['cellsRead'] + 1]),
                                        'minimumRecordCharacters': min(costs), '6000MinimumCharacters': 6000 * min(costs)}
    workbook.close()

    workbook = new()
    workbook.active['A1'] = 'source'
    for count in (18, 19):
        observe('ranges-' + str(count), workbook, 'read', {'ranges': [{'sheet': 'Rules', 'range': 'A1'} for _ in range(count)]},
                {'sameCellRepeated': count})
    for area in (12000, 12001):
        observe('area-' + str(area), workbook, 'read', {'ranges': [{'sheet': 'Rules', 'range': 'A1:A' + str(area)}]},
                {'rectangleArea': area, 'populatedCells': 1})
    workbook.close()

    workbook = new()
    for number in range(1, 51):
        workbook.active.cell(number, 1, 'x' * 4000)
    row = observe('overlapping-ranges', workbook, 'read', {'ranges': [
        {'sheet': 'Rules', 'range': 'A1:A50'}, {'sheet': 'Rules', 'range': 'A1:B50'}]},
        {'populatedCells': 50, 'textCharacters': 4000, 'overlap': 'same 50 records appear in both range arrays'})
    row['responseCharacters'] = len(json.dumps(row['result']['output'], ensure_ascii=False))
    workbook.close()

    workbook = new()
    workbook.active['A1'] = 'x' * 4001
    workbook.active['A2'] = 'commented source'
    workbook.active['A2'].comment = Comment('c' * 4001, 'Synthetic')
    workbook.active['A3'] = '=' + '1+' * 600 + '1'
    observe('record-caps', workbook, 'read', {'ranges': [{'sheet': 'Rules', 'range': 'A1:A3'}]},
            {'A1': {'textLength': 4001}, 'A2': {'commentLength': 4001}, 'A3': {'formulaLength': 1202, 'cachedValue': None}})
    workbook.close()

    # EF-05: actual stored cells at the maximum coordinate, without expanding the grid.
    workbook = new()
    workbook.active['A1'] = 'near-source'
    workbook.active['XFD1048576'] = 'far-source'
    workbook.active['XFC1048576'].fill = PatternFill('solid', fgColor='FFFFFF')
    far_recipe = {'meaningfulCells': {'A1': 'near-source', 'XFD1048576': 'far-source'},
                  'styleOnlyCells': ['XFC1048576'], 'noPerformanceThreshold': True}
    far_inventory = observe('far-cell-inventory', workbook, 'inventory', recipe=far_recipe)
    observe('far-cell-detail', workbook, 'read',
            {'ranges': [{'sheet': 'Rules', 'range': 'A1'}, {'sheet': 'Rules', 'range': 'XFC1048576:XFD1048576'}]},
            far_recipe, existing_input=root / far_inventory['input']['workbook'])
    workbook.close()

    # Physical sheet reorder: identical source cells, only workbook sheet order changes.
    workbook = new()
    workbook.active.title = 'Alpha requirements'
    workbook.create_sheet('Omega requirements')
    source_rows = {'Alpha requirements': ['Flux cohesion', '17 이상', 'MPa', '7일'],
                   'Omega requirements': ['Bending limit', '29 이하', 'kPa', '24시간']}
    for name, values in source_rows.items():
        sheet = workbook[name]
        for number, value in enumerate(['항목', '기준', '단위', '시험조건'], start=1):
            sheet.cell(1, number, value)
        for number, value in enumerate(values, start=1):
            sheet.cell(2, number, value)
    for order in ('forward', 'reverse'):
        if order == 'reverse':
            workbook.move_sheet('Omega requirements', offset=-1)
        recipe = {'sheetOrder': list(workbook.sheetnames), 'sourceRows': source_rows,
                  'transformation': 'public Workbook.move_sheet only' if order == 'reverse' else 'original order'}
        inventory_row = observe('sheet-order-' + order + '-inventory', workbook, 'inventory', recipe=recipe)
        observe('sheet-order-' + order + '-detail', workbook, 'read',
                {'ranges': [{'sheet': name, 'range': 'A1:D2'} for name in workbook.sheetnames]},
                recipe, existing_input=root / inventory_row['input']['workbook'])
    workbook.close()

    projections = []
    for row in rows:
        value = row['result'].get('output', {})
        projections.append({'id': row['id'], 'accepted': row['result']['accepted'], **(
            {'complete': value['complete'], 'sheets': len(value['sheets']),
             'profiledCells': sum(sheet['profiledCells'] for sheet in value['sheets']),
             'regions': sum(len(sheet['regions']) for sheet in value['sheets']),
             'regionsOmitted': sum(sheet['regionsOmitted'] for sheet in value['sheets'])}
            if 'sheets' in value else
            {'complete': value['complete'], 'cellsRead': value['cellsRead']} if 'ranges' in value else {})})
    output = {'runtime': {'python': sys.version, 'openpyxl': openpyxl.__version__}, 'product': {'path': 'server/criteria-workbook-profile.py',
              'sha256': hashlib.sha256(product_path.read_bytes()).hexdigest()}, 'rows': rows, 'projection': projections}
    (destination / 'actual.json').write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'artifact': str(destination / 'actual.json'), 'rows': len(rows)}))


if __name__ == '__main__':
    main()
